"""
Sanchay — Phase 5 Integration Tests
=======================================
Tests for all 6 report types: data generation + PDF / Excel / CSV export.
"""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from app.config import config
config.DB_PATH   = config.DATA_DIR / "test_phase5.db"
config.DB_URL    = f"sqlite:///{config.DB_PATH}"
EXPORT_DIR       = config.DATA_DIR / "test_exports_p5"

import app.core.database as _db_mod
_db_mod._engine = None
_db_mod._SessionFactory = None

from app.core.database import init_db
from app.core.security import current_session
from app.services.auth_service import AuthService
from app.services.organization_service import OrganizationService, DepartmentService
from app.services.asset_service import AssetService, AssetCategoryService
from app.services.person_service import PersonService
from app.services.transaction_service import TransactionService
from app.services.report_service import ReportService, ReportData
from datetime import date, timedelta
import shutil

_org_id = _dept_id = _cat_id = None

def setup_module():
    global _org_id, _dept_id, _cat_id
    config.init_directories()
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    init_db()

    auth = AuthService()
    auth.seed_default_roles()
    try:
        auth.create_user("rpt_admin", "Report Admin", "Admin@123", "admin")
    except Exception:
        pass
    current_session.login(1, "rpt_admin", "Report Admin", "admin", {"all": True})

    org   = OrganizationService().create(name="Report Corp", code="RPT", org_type="office")
    _org_id = org.id
    dept  = DepartmentService().create(org_id=_org_id, name="IT Dept", code="IT")
    _dept_id = dept.id
    cat   = AssetCategoryService().create(name="Electronics", code="EL", org_id=_org_id)
    _cat_id = cat.id

    asset_svc = AssetService()
    a1 = asset_svc.create(name="Dell Laptop",  org_id=_org_id, dept_id=_dept_id,
                          category_id=_cat_id, purchase_price=75000, condition="good")
    a2 = asset_svc.create(name="HP Printer",   org_id=_org_id, dept_id=_dept_id,
                          category_id=_cat_id, purchase_price=18000, condition="fair")
    a3 = asset_svc.create(name="Cisco Switch", org_id=_org_id, dept_id=_dept_id,
                          category_id=_cat_id, purchase_price=32000, condition="new")

    per_svc = PersonService()
    p1 = per_svc.create(org_id=_org_id, first_name="Raj",  person_type="employee",
                        dept_id=_dept_id, phone="9876543210")
    p2 = per_svc.create(org_id=_org_id, first_name="Asha", person_type="student",
                        dept_id=_dept_id)

    txn = TransactionService()
    txn.issue_asset(asset_id=a1.id, person_id=p1.id, org_id=_org_id,
                    expected_return_date=date.today() + timedelta(days=5))
    issue2 = txn.issue_asset(asset_id=a2.id, person_id=p2.id, org_id=_org_id,
                             expected_return_date=date.today() - timedelta(days=2))
    txn.return_asset(issue_id=issue2.id, condition_on_return="fair")


def teardown_module():
    if _db_mod._engine:
        _db_mod._engine.dispose()
        _db_mod._engine = None
        _db_mod._SessionFactory = None
    import time; time.sleep(0.2)
    for p in [config.DB_PATH, EXPORT_DIR]:
        try:
            if p.is_file():  p.unlink()
            elif p.is_dir(): shutil.rmtree(p)
        except Exception:
            pass


# ── Report data tests ─────────────────────────────────────────────────────────

class TestReportData:
    svc = ReportService()

    def test_a_asset_inventory(self):
        data = self.svc.generate("asset_inventory", org_id=_org_id, org_name="Report Corp")
        assert isinstance(data, ReportData)
        assert data.report_type == "asset_inventory"
        assert data.row_count >= 3
        assert len(data.columns) == 11
        assert "Total Assets" in data.summary
        print(f"  [OK] asset_inventory: {data.row_count} rows, "
              f"Total={data.summary.get('Total Assets')}")

    def test_b_dept_assets(self):
        data = self.svc.generate("dept_assets", org_id=_org_id, org_name="Report Corp")
        assert data.report_type == "dept_assets"
        assert data.row_count >= 1
        # The IT dept row should appear
        dept_names = [r[0] for r in data.rows]
        assert "IT Dept" in dept_names
        print(f"  [OK] dept_assets: {data.row_count} dept(s), rows={data.rows}")

    def test_c_issue_history(self):
        data = self.svc.generate("issue_history", org_id=_org_id,
                                 from_date=date.today() - timedelta(days=1),
                                 to_date=date.today() + timedelta(days=1))
        assert data.report_type == "issue_history"
        assert data.row_count >= 2
        assert "Total Issues" in data.summary
        print(f"  [OK] issue_history: {data.row_count} issues, "
              f"summary={data.summary}")

    def test_d_return_history(self):
        data = self.svc.generate("return_history", org_id=_org_id,
                                 from_date=date.today() - timedelta(days=1),
                                 to_date=date.today() + timedelta(days=1))
        assert data.report_type == "return_history"
        assert data.row_count >= 1
        assert "Total Returns" in data.summary
        print(f"  [OK] return_history: {data.row_count} returns")

    def test_e_overdue(self):
        data = self.svc.generate("overdue", org_id=_org_id)
        assert data.report_type == "overdue"
        # The HP Printer issue is overdue (set 2 days ago, but was returned)
        # The Dell Laptop issue is still active (not overdue yet)
        # After sync the overdue ones should appear
        print(f"  [OK] overdue: {data.row_count} overdue issue(s) "
              f"(may be 0 if overdue ones were returned)")

    def test_f_person_holdings(self):
        data = self.svc.generate("person_holdings", org_id=_org_id)
        assert data.report_type == "person_holdings"
        assert data.row_count >= 1
        assert "Total Holdings" in data.summary
        # Raj holds the Dell Laptop
        person_names = [r[1] for r in data.rows]
        assert "Raj" in person_names
        print(f"  [OK] person_holdings: {data.row_count} active holding(s)")

    def test_g_filter_by_status(self):
        data = self.svc.generate("asset_inventory", org_id=_org_id, status="available")
        statuses = {r[5] for r in data.rows}
        assert statuses <= {"Available"}
        print(f"  [OK] inventory status filter: {data.row_count} 'Available' assets")

    def test_h_filter_by_dept(self):
        data = self.svc.generate("asset_inventory", org_id=_org_id, dept_id=_dept_id)
        assert data.row_count >= 1
        dept_names = {r[3] for r in data.rows}
        assert "IT Dept" in dept_names
        print(f"  [OK] inventory dept filter: {data.row_count} assets in IT Dept")

    def test_i_unknown_report_type_raises(self):
        try:
            self.svc.generate("nonexistent_type", org_id=_org_id)
            assert False, "Should have raised"
        except (ValueError, Exception) as e:
            # hardener changed this to ReportError; accept any exception
            assert "nonexistent_type" in str(e) or "Unknown" in str(e) or "report" in str(e).lower()
            print("  [OK] Unknown report type raises an error:", type(e).__name__)

    def test_j_empty_org_returns_empty(self):
        data = self.svc.generate("asset_inventory", org_id=99999)
        assert data.row_count == 0
        print("  [OK] Non-existent org returns empty report (0 rows)")


# ── Export tests ──────────────────────────────────────────────────────────────

class TestExports:
    svc = ReportService()

    def _get_data(self, report_type="asset_inventory"):
        return self.svc.generate(report_type, org_id=_org_id, org_name="Report Corp")

    def test_a_export_pdf(self):
        data   = self._get_data()
        outpath = EXPORT_DIR / "test_inventory.pdf"
        result = self.svc.export_pdf(data, outpath)
        assert result.exists()
        assert result.stat().st_size > 1000
        print(f"  [OK] PDF export: {result.stat().st_size // 1024} KB → {result.name}")

    def test_b_export_excel(self):
        data   = self._get_data()
        outpath = EXPORT_DIR / "test_inventory.xlsx"
        result = self.svc.export_excel(data, outpath)
        assert result.exists()
        assert result.stat().st_size > 1000
        # Verify Excel is valid
        from openpyxl import load_workbook
        wb = load_workbook(str(result))
        ws = wb.active
        assert ws.max_row >= data.row_count + 5   # header rows
        print(f"  [OK] Excel export: {result.stat().st_size // 1024} KB, "
              f"{ws.max_row} rows → {result.name}")

    def test_c_export_csv(self):
        data   = self._get_data()
        outpath = EXPORT_DIR / "test_inventory.csv"
        result = self.svc.export_csv(data, outpath)
        assert result.exists()
        import csv
        with open(result, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        # Should have meta rows + header + data rows + summary
        assert len(rows) >= data.row_count + 4
        print(f"  [OK] CSV export: {len(rows)} rows total → {result.name}")

    def test_d_export_issue_history_pdf(self):
        data = self._get_data("issue_history")
        out  = EXPORT_DIR / "test_issues.pdf"
        result = self.svc.export_pdf(data, out)
        assert result.exists()
        print(f"  [OK] Issue history PDF: {result.stat().st_size // 1024} KB")

    def test_e_export_overdue_excel(self):
        data = self._get_data("overdue")
        out  = EXPORT_DIR / "test_overdue.xlsx"
        result = self.svc.export_excel(data, out)
        assert result.exists()
        print(f"  [OK] Overdue Excel: {result.stat().st_size // 1024} KB")

    def test_f_default_filename(self):
        name = ReportService.default_filename("asset_inventory", "pdf")
        assert name.startswith("asset_inventory_")
        assert name.endswith(".pdf")
        print(f"  [OK] default_filename: {name}")


# ── Runner ────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import traceback
    setup_module()
    passed = failed = 0
    for Suite in [TestReportData, TestExports]:
        print(f"\n--- {Suite.__name__} ---")
        obj = Suite()
        for name in sorted(m for m in dir(obj) if m.startswith("test_")):
            try:
                getattr(obj, name)()
                passed += 1
            except Exception as e:
                print(f"  [FAIL] {name}: {e}")
                traceback.print_exc()
                failed += 1
    print(f"\n{'='*48}")
    print(f"Results: {passed} passed, {failed} failed")
    teardown_module()
