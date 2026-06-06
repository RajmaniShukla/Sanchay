"""
Sanchay — Report Service
==========================
Data gathering for all six report types.
Produces a ReportData object consumed by PDF / Excel / CSV exporters.
"""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass, field
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path
from typing import Optional

from loguru import logger

from app.config import config
from app.core.database import get_db
from app.core.exceptions import ReportError
from app.models.asset import Asset
from app.models.transaction import AssetIssue, AssetReturn
from app.models.person import Person


# ── Report data container ────────────────────────────────────────────────────

@dataclass
class ReportData:
    """All data needed to render / export a report."""
    report_type:  str
    title:        str
    subtitle:     str
    columns:      list[str]
    rows:         list[list[str]]          # every cell is a plain string
    summary:      dict[str, str] = field(default_factory=dict)
    generated_at: datetime       = field(default_factory=datetime.now)
    org_name:     str            = "Sanchay"
    filters:      dict[str, str] = field(default_factory=dict)

    @property
    def row_count(self) -> int:
        """Return the number of data rows."""
        return len(self.rows)


# ── Report definitions ────────────────────────────────────────────────────────

REPORT_TYPES: dict[str, str] = {
    "asset_inventory":  "Asset Inventory",
    "dept_assets":      "Department-wise Assets",
    "issue_history":    "Issue History",
    "return_history":   "Return History",
    "overdue":          "Overdue Assets",
    "person_holdings":  "Person Asset Holdings",
}

# Rows above this limit are truncated in PDF only (CSV/Excel export all rows)
PDF_ROW_LIMIT = 500


def _fmt_price(val) -> str:
    """Format a numeric value as a price string."""
    if val is None:
        return "—"
    try:
        return f"Rs. {float(val):,.2f}"
    except Exception:
        return str(val)


def _fmt_date(val) -> str:
    """Format a date or datetime as a human-readable string."""
    if val is None:
        return "—"
    if isinstance(val, (date, datetime)):
        return val.strftime("%d %b %Y")
    return str(val)


# ── Service class ─────────────────────────────────────────────────────────────

class ReportService:
    """
    Generates ReportData for any report type.
    Also handles PDF, Excel, and CSV export.
    """

    # ── Class-level report type registry ─────────────────────────────────────

    @classmethod
    def get_report_types(cls) -> dict:
        """Return the mapping of report_type keys to human-readable names."""
        return dict(REPORT_TYPES)

    # ── Generate ─────────────────────────────────────────────────────────────

    def generate(
        self,
        report_type: str,
        org_id: Optional[int] = None,
        org_name: str = "Organization",
        dept_id: Optional[int] = None,
        category_id: Optional[int] = None,
        status: Optional[str] = None,
        person_type: Optional[str] = None,
        from_date: Optional[date] = None,
        to_date: Optional[date] = None,
    ) -> ReportData:
        """
        Generate report data for the given type and filters.

        Args:
            report_type : one of REPORT_TYPES keys
            org_id      : organisation filter (None → returns empty dataset)
            org_name    : display name used in headers
            ...         : other filters (optional)

        Returns:
            ReportData ready to render or export.

        Raises:
            ReportError: If report generation fails due to DB or logic errors.
        """
        generators = {
            "asset_inventory": self._asset_inventory,
            "dept_assets":     self._dept_assets,
            "issue_history":   self._issue_history,
            "return_history":  self._return_history,
            "overdue":         self._overdue,
            "person_holdings": self._person_holdings,
        }

        # Validate report type FIRST — before any org_id checks
        fn = generators.get(report_type)
        if not fn:
            raise ReportError(
                f"Unknown report type '{report_type}'. "
                f"Valid types: {', '.join(REPORT_TYPES.keys())}"
            )

        # Defensive: if no org_id is selected, return an empty placeholder
        # report instead of letting individual generators hit the DB.
        if org_id is None:
            title = REPORT_TYPES.get(report_type, report_type)
            return ReportData(
                report_type=report_type,
                title=title,
                subtitle="No organization selected. Please log in to an organization.",
                columns=[],
                rows=[],
                summary={"Note": "No organisation selected."},
                org_name=org_name,
            )

        try:
            data = fn(
                org_id=org_id,
                dept_id=dept_id,
                category_id=category_id,
                status=status,
                person_type=person_type,
                from_date=from_date,
                to_date=to_date,
            )
        except ReportError:
            raise
        except Exception as exc:
            logger.error(f"Report '{report_type}' failed: {exc}")
            raise ReportError(
                f"Could not generate the '{REPORT_TYPES.get(report_type, report_type)}' report. "
                "Please try again or contact your system administrator."
            ) from exc

        data.org_name = org_name
        logger.info(f"Report '{report_type}' generated: {data.row_count} rows")
        return data

    # ── Report 1: Asset Inventory ─────────────────────────────────────────────

    def _asset_inventory(self, org_id, dept_id, category_id, status, **_) -> ReportData:
        """Generate asset inventory report data."""
        if org_id is None:
            return ReportData(
                report_type="asset_inventory",
                title="Asset Inventory Report",
                subtitle="Please select an organisation to view its asset inventory.",
                columns=["Asset Code", "Name", "Category", "Department",
                         "Serial No.", "Status", "Condition",
                         "Purchase Date", "Purchase Price", "Warranty Expiry", "Location"],
                rows=[],
                summary={"Note": "No organisation selected."},
            )

        from app.repositories.asset_repository import AssetRepository

        try:
            with get_db() as db:
                repo   = AssetRepository(db)
                assets, total = repo.search(
                    query="", org_id=org_id, status=status,
                    category_id=category_id, dept_id=dept_id, limit=5000,
                )

                rows = []
                total_value = 0.0
                for a in assets:
                    price = float(a.purchase_price) if a.purchase_price else 0
                    total_value += price
                    rows.append([
                        a.asset_code,
                        a.name,
                        a.category.name if a.category else "—",
                        a.department.name if a.department else "—",
                        a.serial_number or "—",
                        a.status.capitalize(),
                        (a.condition or "—").capitalize(),
                        _fmt_date(a.purchase_date),
                        _fmt_price(a.purchase_price),
                        _fmt_date(a.warranty_expiry_date),
                        a.location or "—",
                    ])
        except Exception as exc:
            raise ReportError(f"Failed to load asset inventory data: {exc}") from exc

        filters = {}
        if status:     filters["Status"]     = status.capitalize()
        if dept_id:    filters["Department"] = f"ID {dept_id}"
        if category_id: filters["Category"] = f"ID {category_id}"

        return ReportData(
            report_type="asset_inventory",
            title="Asset Inventory Report",
            subtitle=f"All assets in the system as of {_fmt_date(date.today())}",
            columns=[
                "Asset Code", "Name", "Category", "Department",
                "Serial No.", "Status", "Condition",
                "Purchase Date", "Purchase Price", "Warranty Expiry", "Location",
            ],
            rows=rows,
            summary={
                "Total Assets":    str(total),
                "Total Value":     _fmt_price(total_value),
                "Available":       str(sum(1 for a in assets if a.status == "available")),
                "Issued":          str(sum(1 for a in assets if a.status == "issued")),
                "Maintenance":     str(sum(1 for a in assets if a.status == "maintenance")),
                "Disposed / Lost": str(sum(1 for a in assets if a.status in ("disposed", "lost"))),
            },
            filters=filters,
        )

    # ── Report 2: Department-wise Assets ──────────────────────────────────────

    def _dept_assets(self, org_id, dept_id, **_) -> ReportData:
        """Generate department-wise asset distribution report data."""
        if org_id is None and dept_id is None:
            return ReportData(
                report_type="dept_assets",
                title="Department-wise Asset Report",
                subtitle="Please select an organisation to view department asset distribution.",
                columns=["Department", "Code", "Total Assets",
                         "Available", "Issued", "Other", "Total Value"],
                rows=[],
                summary={"Note": "No organisation selected."},
            )

        from app.repositories.organization_repository import DepartmentRepository
        from app.repositories.asset_repository import AssetRepository

        try:
            with get_db() as db:
                dept_repo  = DepartmentRepository(db)
                asset_repo = AssetRepository(db)

                if dept_id:
                    depts = [dept_repo.get_by_id(dept_id)]
                else:
                    depts = dept_repo.get_by_org(org_id) if org_id else []

                rows = []
                grand_total = 0
                for dept in depts:
                    if not dept:
                        continue
                    assets, _ = asset_repo.search(query="", dept_id=dept.id, limit=5000)
                    available  = sum(1 for a in assets if a.status == "available")
                    issued     = sum(1 for a in assets if a.status == "issued")
                    total_val  = sum(float(a.purchase_price or 0) for a in assets)
                    grand_total += len(assets)
                    rows.append([
                        dept.name,
                        dept.code,
                        str(len(assets)),
                        str(available),
                        str(issued),
                        str(len(assets) - available - issued),
                        _fmt_price(total_val),
                    ])
        except Exception as exc:
            raise ReportError(f"Failed to load department asset data: {exc}") from exc

        filters = {}
        if dept_id: filters["Department"] = f"ID {dept_id}"

        return ReportData(
            report_type="dept_assets",
            title="Department-wise Asset Report",
            subtitle=f"Asset distribution across departments as of {_fmt_date(date.today())}",
            columns=[
                "Department", "Code", "Total Assets",
                "Available", "Issued", "Other", "Total Value",
            ],
            rows=rows,
            summary={"Grand Total Assets": str(grand_total)},
            filters=filters,
        )

    # ── Report 3: Issue History ────────────────────────────────────────────────

    def _issue_history(self, org_id, status, from_date, to_date, **_) -> ReportData:
        """Generate asset issue history report data."""
        try:
            with get_db() as db:
                from app.repositories.transaction_repository import TransactionRepository
                repo = TransactionRepository(db)
                issues, total = repo.get_all_issues(
                    org_id=org_id, status=status,
                    from_date=from_date, to_date=to_date, limit=5000,
                )

                rows = []
                for issue in issues:
                    asset  = issue.asset
                    person = issue.person
                    ret    = issue.returns[0] if issue.returns else None
                    rows.append([
                        f"#{issue.id}",
                        asset.asset_code if asset else "—",
                        asset.name       if asset else "—",
                        person.full_name if person else "—",
                        person.person_type.capitalize() if person else "—",
                        issue.issued_by.username if issue.issued_by else "—",
                        _fmt_date(issue.issue_date),
                        _fmt_date(issue.expected_return_date),
                        _fmt_date(ret.return_date) if ret else "—",
                        issue.status.capitalize(),
                        issue.purpose or "—",
                    ])
        except Exception as exc:
            raise ReportError(f"Failed to load issue history data: {exc}") from exc

        filters = {}
        if from_date: filters["From"]   = _fmt_date(from_date)
        if to_date:   filters["To"]     = _fmt_date(to_date)
        if status:    filters["Status"] = status.capitalize()

        return ReportData(
            report_type="issue_history",
            title="Asset Issue History Report",
            subtitle="Complete record of asset issues",
            columns=[
                "Issue #", "Asset Code", "Asset Name",
                "Issued To", "Type", "Issued By",
                "Issue Date", "Expected Return", "Return Date",
                "Status", "Purpose",
            ],
            rows=rows,
            summary={
                "Total Issues":   str(total),
                "Active":         str(sum(1 for r in rows if r[9] == "Active")),
                "Returned":       str(sum(1 for r in rows if r[9] == "Returned")),
                "Overdue":        str(sum(1 for r in rows if r[9] == "Overdue")),
            },
            filters=filters,
        )

    # ── Report 4: Return History ──────────────────────────────────────────────

    def _return_history(self, org_id, from_date, to_date, **_) -> ReportData:
        """Generate asset return history report data."""
        try:
            with get_db() as db:
                from app.repositories.transaction_repository import TransactionRepository
                repo = TransactionRepository(db)
                returns, total = repo.get_all_returns(
                    org_id=org_id, from_date=from_date, to_date=to_date, limit=5000,
                )

                rows = []
                for ret in returns:
                    asset  = ret.asset
                    person = ret.person
                    recvd  = ret.returned_to
                    rows.append([
                        f"#{ret.id}",
                        f"#{ret.issue_id}",
                        asset.asset_code if asset else "—",
                        asset.name       if asset else "—",
                        person.full_name if person else "—",
                        recvd.username   if recvd  else "—",
                        _fmt_date(ret.return_date),
                        (ret.condition_on_return or "—").capitalize(),
                        ret.remarks or "—",
                    ])
        except Exception as exc:
            raise ReportError(f"Failed to load return history data: {exc}") from exc

        filters = {}
        if from_date: filters["From"] = _fmt_date(from_date)
        if to_date:   filters["To"]   = _fmt_date(to_date)

        return ReportData(
            report_type="return_history",
            title="Asset Return History Report",
            subtitle="Complete record of asset returns",
            columns=[
                "Return #", "Issue #", "Asset Code", "Asset Name",
                "Returned By", "Received By",
                "Return Date", "Condition", "Remarks",
            ],
            rows=rows,
            summary={"Total Returns": str(total)},
            filters=filters,
        )

    # ── Report 5: Overdue Assets ──────────────────────────────────────────────

    def _overdue(self, org_id, **_) -> ReportData:
        """Generate overdue assets report data."""
        today = date.today()

        try:
            with get_db() as db:
                from app.repositories.transaction_repository import TransactionRepository
                repo   = TransactionRepository(db)
                issues = repo.get_overdue_issues(org_id)

                rows = []
                for issue in issues:
                    asset  = issue.asset
                    person = issue.person
                    exp    = issue.expected_return_date
                    days_late = (today - exp).days if exp else 0
                    rows.append([
                        f"#{issue.id}",
                        asset.asset_code if asset else "—",
                        asset.name       if asset else "—",
                        person.full_name if person else "—",
                        person.person_type.capitalize() if person else "—",
                        person.phone     if person else "—",
                        _fmt_date(issue.issue_date),
                        _fmt_date(exp),
                        str(days_late),
                        issue.purpose or "—",
                    ])
        except Exception as exc:
            raise ReportError(f"Failed to load overdue assets data: {exc}") from exc

        return ReportData(
            report_type="overdue",
            title="Overdue Assets Report",
            subtitle=f"Assets not returned by expected date — as of {_fmt_date(today)}",
            columns=[
                "Issue #", "Asset Code", "Asset Name",
                "Issued To", "Type", "Contact",
                "Issue Date", "Expected Return", "Days Overdue", "Purpose",
            ],
            rows=rows,
            summary={"Total Overdue": str(len(rows))},
            filters={},
        )

    # ── Report 6: Person Holdings ─────────────────────────────────────────────

    def _person_holdings(self, org_id, dept_id, person_type, **_) -> ReportData:
        """Generate person asset holdings report data."""
        if org_id is None:
            return ReportData(
                report_type="person_holdings",
                title="Person Asset Holdings Report",
                subtitle="Please select an organisation to view person holdings.",
                columns=["Person ID", "Name", "Type", "Department", "Contact",
                         "Asset Code", "Asset Name", "Category",
                         "Issue Date", "Expected Return"],
                rows=[],
                summary={"Note": "No organisation selected."},
            )

        try:
            with get_db() as db:
                from app.repositories.transaction_repository import TransactionRepository
                from app.repositories.person_repository import PersonRepository

                p_repo   = PersonRepository(db)
                txn_repo = TransactionRepository(db)

                persons = p_repo.get_active_by_org(org_id) if org_id else []
                if person_type:
                    persons = [p for p in persons if p.person_type == person_type]
                if dept_id:
                    persons = [p for p in persons if p.dept_id == dept_id]

                rows = []
                for person in persons:
                    issues = txn_repo.get_issues_for_person(person.id, status="active")
                    if not issues:
                        continue
                    for issue in issues:
                        asset = issue.asset
                        rows.append([
                            person.display_code,
                            person.full_name,
                            person.person_type.capitalize(),
                            person.department.name if person.department else "—",
                            person.phone or "—",
                            asset.asset_code if asset else "—",
                            asset.name       if asset else "—",
                            asset.category.name if (asset and asset.category) else "—",
                            _fmt_date(issue.issue_date),
                            _fmt_date(issue.expected_return_date),
                        ])
        except Exception as exc:
            raise ReportError(f"Failed to load person holdings data: {exc}") from exc

        filters = {}
        if person_type: filters["Person Type"] = person_type.capitalize()
        if dept_id:     filters["Department"]  = f"ID {dept_id}"

        return ReportData(
            report_type="person_holdings",
            title="Person Asset Holdings Report",
            subtitle=f"Active asset holdings per person as of {_fmt_date(date.today())}",
            columns=[
                "Person ID", "Name", "Type", "Department", "Contact",
                "Asset Code", "Asset Name", "Category",
                "Issue Date", "Expected Return",
            ],
            rows=rows,
            summary={"Total Holdings": str(len(rows))},
            filters=filters,
        )

    # ── Exporters ─────────────────────────────────────────────────────────────

    def export_pdf(self, data: ReportData, filepath: Path) -> Path:
        """Export ReportData to a styled PDF using ReportLab."""
        # Defensive: ensure parent directory exists
        filepath = Path(filepath)
        try:
            filepath.parent.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            raise ReportError(
                f"Cannot create directory '{filepath.parent}'. Check write permissions."
            ) from e

        # Defensive: verify ReportLab is available
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm, mm
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle,
                Paragraph, Spacer, HRFlowable,
            )
            from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        except ImportError as e:
            raise ReportError(
                "PDF generation failed. Check that ReportLab is installed. "
                "Run: pip install reportlab"
            ) from e

        # Truncate for PDF if over limit
        pdf_rows = data.rows
        truncated = False
        if len(pdf_rows) > PDF_ROW_LIMIT:
            pdf_rows = pdf_rows[:PDF_ROW_LIMIT]
            truncated = True
            logger.info(
                f"PDF export: report truncated from {data.row_count} to {PDF_ROW_LIMIT} rows."
            )

        # ── Use landscape for wide reports ─────────────────────────────────────
        page_size = landscape(A4) if len(data.columns) > 7 else A4
        page_w, page_h = page_size
        margin = 1.5 * cm

        doc = SimpleDocTemplate(
            str(filepath),
            pagesize=page_size,
            leftMargin=margin, rightMargin=margin,
            topMargin=margin, bottomMargin=margin + 0.6 * cm,
            title=data.title,
            author=config.APP_NAME,
        )

        styles = getSampleStyleSheet()
        PRIMARY   = colors.HexColor("#2563EB")
        LIGHT_BG  = colors.HexColor("#EFF6FF")
        ALT_ROW   = colors.HexColor("#F8FAFC")
        BORDER    = colors.HexColor("#E2E8F0")
        DARK_TEXT = colors.HexColor("#0F172A")
        MUTED     = colors.HexColor("#64748B")

        h1_style = ParagraphStyle("h1", parent=styles["Heading1"],
            fontSize=18, textColor=PRIMARY, spaceAfter=2, spaceBefore=0,
            fontName="Helvetica-Bold")
        h2_style = ParagraphStyle("h2", parent=styles["Normal"],
            fontSize=10, textColor=MUTED, spaceAfter=4)
        label_style = ParagraphStyle("lbl", parent=styles["Normal"],
            fontSize=9, textColor=MUTED, fontName="Helvetica-Bold")
        val_style = ParagraphStyle("val", parent=styles["Normal"],
            fontSize=9, textColor=DARK_TEXT)
        cell_style = ParagraphStyle("cell", parent=styles["Normal"],
            fontSize=8, textColor=DARK_TEXT, wordWrap="LTR")
        hdr_style = ParagraphStyle("hdr", parent=styles["Normal"],
            fontSize=8, textColor=colors.white, fontName="Helvetica-Bold",
            alignment=TA_CENTER)

        story = []

        # ── Page header ────────────────────────────────────────────────────────
        story.append(Paragraph(data.org_name, h2_style))
        story.append(Paragraph(data.title, h1_style))
        story.append(Paragraph(data.subtitle, h2_style))

        if truncated:
            truncation_note = ParagraphStyle("trunc", parent=styles["Normal"],
                fontSize=9, textColor=colors.HexColor("#D97706"), fontName="Helvetica-Bold")
            story.append(Paragraph(
                f"⚠ Report truncated to {PDF_ROW_LIMIT} rows for PDF. "
                "Export to CSV or Excel to get all rows.",
                truncation_note,
            ))

        story.append(HRFlowable(width="100%", thickness=2, color=PRIMARY, spaceAfter=6))

        # Meta row: generated date + filters
        meta_parts = [f"Generated: {data.generated_at.strftime('%d %b %Y  %H:%M')}"]
        for k, v in data.filters.items():
            meta_parts.append(f"  {k}: {v}")
        story.append(Paragraph("  |  ".join(meta_parts),
                                ParagraphStyle("meta", parent=styles["Normal"],
                                               fontSize=8, textColor=MUTED)))
        story.append(Spacer(1, 8))

        # ── Data table ─────────────────────────────────────────────────────────
        if pdf_rows:
            avail_w = page_w - 2 * margin
            col_w   = avail_w / len(data.columns)
            col_widths = [col_w] * len(data.columns)

            # Header row
            hdr_row = [Paragraph(c, hdr_style) for c in data.columns]
            table_data = [hdr_row]

            for row_data in pdf_rows:
                table_data.append([
                    Paragraph(str(v) if v else "—", cell_style)
                    for v in row_data
                ])

            tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
            tbl.setStyle(TableStyle([
                # Header
                ("BACKGROUND",  (0, 0), (-1, 0),  PRIMARY),
                ("TEXTCOLOR",   (0, 0), (-1, 0),  colors.white),
                ("FONTNAME",    (0, 0), (-1, 0),  "Helvetica-Bold"),
                ("FONTSIZE",    (0, 0), (-1, 0),  8),
                ("ALIGN",       (0, 0), (-1, 0),  "CENTER"),
                ("TOPPADDING",  (0, 0), (-1, 0),  6),
                ("BOTTOMPADDING",(0, 0),(-1, 0),  6),
                # Data rows
                ("FONTSIZE",    (0, 1), (-1, -1), 8),
                ("TOPPADDING",  (0, 1), (-1, -1), 4),
                ("BOTTOMPADDING",(0, 1),(-1, -1), 4),
                ("ALIGN",       (0, 1), (-1, -1), "LEFT"),
                ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
                # Alternating rows
                *[("BACKGROUND", (0, r), (-1, r), ALT_ROW)
                  for r in range(2, len(table_data), 2)],
                # Grid
                ("GRID",      (0, 0), (-1, -1), 0.5, BORDER),
                ("LINEBELOW", (0, 0), (-1, 0),  1.5, PRIMARY),
            ]))
            story.append(tbl)
        else:
            story.append(Paragraph("No data found for the selected filters.",
                                   ParagraphStyle("empty", parent=styles["Normal"],
                                                  fontSize=11, textColor=MUTED,
                                                  alignment=TA_CENTER)))

        story.append(Spacer(1, 12))

        # ── Summary box ────────────────────────────────────────────────────────
        if data.summary:
            sum_data = [
                [Paragraph(k, label_style), Paragraph(v, val_style)]
                for k, v in data.summary.items()
            ]
            sum_tbl = Table(sum_data, colWidths=[5 * cm, 4 * cm])
            sum_tbl.setStyle(TableStyle([
                ("BACKGROUND",     (0, 0), (-1, -1), LIGHT_BG),
                ("BOX",            (0, 0), (-1, -1), 1, PRIMARY),
                ("TOPPADDING",     (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING",  (0, 0), (-1, -1), 4),
                ("LEFTPADDING",    (0, 0), (-1, -1), 8),
            ]))
            story.append(Paragraph("Summary", ParagraphStyle("sumhdr",
                parent=styles["Normal"], fontSize=10, textColor=PRIMARY,
                fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=4)))
            story.append(sum_tbl)

        # ── Page footer via canvas callback ────────────────────────────────────
        def _footer(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 7)
            canvas.setFillColor(MUTED)
            footer_text = (
                f"{config.APP_NAME} v{config.APP_VERSION}  |  "
                f"{data.title}  |  Page {doc.page}"
            )
            canvas.drawString(margin, 0.5 * cm, footer_text)
            canvas.restoreState()

        try:
            doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
        except PermissionError as e:
            raise ReportError(
                f"Cannot write to '{filepath}'. Check write permissions."
            ) from e
        except Exception as e:
            raise ReportError(f"PDF generation failed: {e}") from e
        logger.info(f"PDF exported: {filepath} ({filepath.stat().st_size // 1024} KB)")
        return filepath

    def export_excel(self, data: ReportData, filepath: Path) -> Path:
        """Export ReportData to a styled Excel workbook using openpyxl (no row limit)."""
        # Defensive: ensure parent directory exists
        filepath = Path(filepath)
        try:
            filepath.parent.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            raise ReportError(
                f"Cannot create directory '{filepath.parent}'. Check write permissions."
            ) from e

        # Defensive: verify openpyxl is available
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (
                PatternFill, Font, Alignment, Border, Side, GradientFill,
            )
            from openpyxl.utils import get_column_letter
        except ImportError as e:
            raise ReportError(
                "Excel generation failed. Check that openpyxl is installed. "
                "Run: pip install openpyxl"
            ) from e

        wb = Workbook()
        ws = wb.active
        ws.title = data.title[:31]  # Excel sheet name limit

        PRIMARY_HEX = "2563EB"
        LIGHT_HEX   = "EFF6FF"
        ALT_HEX     = "F8FAFC"
        HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        TITLE_FONT  = Font(name="Calibri", bold=True, color="0F172A", size=14)
        META_FONT   = Font(name="Calibri", italic=True, color="64748B", size=9)
        DATA_FONT   = Font(name="Calibri", size=9)
        BOLD_FONT   = Font(name="Calibri", bold=True, size=9)
        CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
        LEFT        = Alignment(horizontal="left",   vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        row = 1

        # ── Title rows ─────────────────────────────────────────────────────────
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=len(data.columns))
        title_cell = ws.cell(row=row, column=1, value=data.org_name)
        title_cell.font = Font(name="Calibri", bold=True, color=PRIMARY_HEX, size=11)
        title_cell.fill = PatternFill("solid", fgColor=LIGHT_HEX)
        title_cell.alignment = LEFT
        ws.row_dimensions[row].height = 20
        row += 1

        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=len(data.columns))
        ws.cell(row=row, column=1, value=data.title).font = TITLE_FONT
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT_HEX)
        ws.cell(row=row, column=1).alignment = LEFT
        ws.row_dimensions[row].height = 22
        row += 1

        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=len(data.columns))
        ws.cell(row=row, column=1,
                value=f"{data.subtitle}  |  Generated: "
                      f"{data.generated_at.strftime('%d %b %Y  %H:%M')}").font = META_FONT
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT_HEX)
        ws.row_dimensions[row].height = 16
        row += 1

        if data.filters:
            flt_text = "  |  ".join(f"{k}: {v}" for k, v in data.filters.items())
            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row, end_column=len(data.columns))
            ws.cell(row=row, column=1, value=f"Filters: {flt_text}").font = META_FONT
            ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT_HEX)
            row += 1

        row += 1  # blank spacer

        # ── Column headers ─────────────────────────────────────────────────────
        for col_idx, col_name in enumerate(data.columns, start=1):
            cell = ws.cell(row=row, column=col_idx, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = PatternFill("solid", fgColor=PRIMARY_HEX)
            cell.alignment = CENTER
            cell.border = thin_border
        ws.row_dimensions[row].height = 22
        data_start_row = row
        row += 1

        # ── Data rows (all rows — no truncation for Excel) ─────────────────────
        for r_idx, row_data in enumerate(data.rows):
            fill_color = LIGHT_HEX if r_idx % 2 == 0 else ALT_HEX
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row, column=col_idx, value=str(val) if val else "—")
                cell.font  = DATA_FONT
                cell.fill  = PatternFill("solid", fgColor=fill_color)
                cell.alignment = LEFT
                cell.border = thin_border
            ws.row_dimensions[row].height = 16
            row += 1

        # ── Summary rows ────────────────────────────────────────────────────────
        if data.summary:
            row += 1
            ws.cell(row=row, column=1, value="Summary").font = Font(
                name="Calibri", bold=True, color=PRIMARY_HEX, size=10)
            row += 1
            for k, v in data.summary.items():
                ws.cell(row=row, column=1, value=k).font = BOLD_FONT
                ws.cell(row=row, column=2, value=v).font  = DATA_FONT
                row += 1

        # ── Column widths ──────────────────────────────────────────────────────
        for col_idx, col_name in enumerate(data.columns, start=1):
            max_len = len(col_name)
            for r in data.rows:
                val = str(r[col_idx - 1]) if col_idx - 1 < len(r) else ""
                max_len = max(max_len, min(len(val), 40))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

        # ── Freeze pane at header ──────────────────────────────────────────────
        ws.freeze_panes = ws.cell(row=data_start_row + 1, column=1)

        # Defensive: validate row cell counts match column count
        col_count = len(data.columns)
        for i, row_data in enumerate(data.rows):
            if len(row_data) != col_count:
                logger.warning(
                    f"Excel export: row {i} has {len(row_data)} cells but "
                    f"{col_count} columns expected; row may render incorrectly."
                )

        try:
            wb.save(str(filepath))
        except PermissionError as e:
            raise ReportError(
                f"Cannot write to '{filepath}'. Check write permissions."
            ) from e
        except Exception as e:
            raise ReportError(f"Excel generation failed: {e}") from e
        logger.info(f"Excel exported: {filepath} ({data.row_count} rows)")
        return filepath

    def export_csv(self, data: ReportData, filepath: Path) -> Path:
        """Export ReportData to a plain CSV file (all rows — no truncation).
        
        Uses UTF-8 BOM encoding (utf-8-sig) for correct display in Excel on Windows.
        """
        filepath = Path(filepath)
        # Defensive: ensure parent directory exists
        try:
            filepath.parent.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            raise ReportError(
                f"Cannot create directory '{filepath.parent}'. Check write permissions."
            ) from e

        try:
            with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                # Meta header
                writer.writerow([data.title])
                writer.writerow([data.subtitle])
                writer.writerow([
                    f"Generated: {data.generated_at.strftime('%d %b %Y %H:%M')}"
                ])
                if data.filters:
                    writer.writerow([
                        "Filters: " + "  |  ".join(f"{k}: {v}"
                        for k, v in data.filters.items())
                    ])
                writer.writerow([])
                # Column headers
                writer.writerow(data.columns)
                # Data — all rows
                for row in data.rows:
                    writer.writerow(row)
                # Summary
                if data.summary:
                    writer.writerow([])
                    writer.writerow(["Summary"])
                    for k, v in data.summary.items():
                        writer.writerow([k, v])
        except PermissionError as e:
            raise ReportError(
                f"Cannot write to '{filepath}'. Check write permissions."
            ) from e
        except Exception as e:
            raise ReportError(f"CSV generation failed: {e}") from e

        logger.info(f"CSV exported: {filepath} ({data.row_count} rows)")
        return filepath

    # ── Default filename helper ────────────────────────────────────────────────

    @staticmethod
    def default_filename(report_type: str, extension: str) -> str:
        """Generate a timestamped default filename for export files."""
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{report_type}_{ts}.{extension}"
